#!/usr/bin/env python3
"""Give the built workbook its cached formula values without letting LibreOffice
rewrite the file we ship.

LibreOffice is the only formula engine available here, but saving through it
rewrites xl/styles.xml: it renames fonts it cannot find on this Linux box to
WenQuanYi Zen Hei (a face the user's Excel does not have), re-serialises floats
at 15 significant digits and nudges a few row heights and number formats.

So we recalculate a throwaway COPY and transplant only the computed values back
into the pristine build, editing each sheet's XML in place so every byte of
styling stays exactly as openpyxl wrote it.
"""
import datetime
import re
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

from openpyxl import load_workbook
from openpyxl.utils.datetime import CALENDAR_WINDOWS_1900 as EPOCH, to_excel

BUILT = Path(sys.argv[1] if len(sys.argv) > 1 else
             "/home/user/Easy-Hard/EasyHardMoney_NDX_2026R4.5.3_COMBINED_ABC.xlsx")
SCRATCH = Path("/tmp/claude-0/-home-user-Easy-Hard/217becff-8454-5f9a-a065-0284fc958152/scratchpad")
RECALC = "/root/.claude/skills/xlsx/scripts/recalc.py"
WORK = SCRATCH / "recalc_donor.xlsx"

ERRORS = {"#DIV/0!", "#N/A", "#NAME?", "#NULL!", "#NUM!", "#REF!", "#VALUE!", "#GETTING_DATA"}


def sheet_parts(path):
    """sheet title -> worksheets/sheetN.xml part name, via the workbook rels."""
    with zipfile.ZipFile(path) as z:
        wb = z.read("xl/workbook.xml").decode("utf8")
        rels = z.read("xl/_rels/workbook.xml.rels").decode("utf8")
    rid_to_target = {}
    for tag in re.findall(r"<Relationship\b[^>]*/?>", rels):
        rid = re.search(r'Id="([^"]+)"', tag)
        target = re.search(r'Target="([^"]+)"', tag)
        if rid and target:
            rid_to_target[rid.group(1)] = target.group(1)
    out = {}
    for m in re.finditer(r'<sheet\b[^>]*/>', wb):
        tag = m.group(0)
        name = re.search(r'name="([^"]*)"', tag).group(1)
        rid = re.search(r'r:id="([^"]+)"', tag).group(1)
        target = rid_to_target[rid].lstrip("/")
        if not target.startswith("xl/"):
            target = "xl/" + target
        out[unescape_attr(name)] = target
    return out


def unescape_attr(s):
    return (s.replace("&amp;", "&").replace("&lt;", "<")
             .replace("&gt;", ">").replace("&quot;", '"').replace("&apos;", "'"))


def cached_values(path):
    """(sheet, coord) -> computed value, harvested from the recalculated donor."""
    wb = load_workbook(path, data_only=True)
    vals = {}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None:
                    vals[(ws.title, c.coordinate)] = c.value
    wb.close()
    return vals


def render(v):
    """Value -> (t attribute or None, <v> text)."""
    if isinstance(v, bool):
        return "b", "1" if v else "0"
    if isinstance(v, (datetime.datetime, datetime.date, datetime.time)):
        # A date-valued formula caches the serial number, not a formatted string;
        # the cell's number format is what turns it back into a date.
        return None, repr(float(to_excel(v, EPOCH)))
    if isinstance(v, datetime.timedelta):
        return None, repr(v.total_seconds() / 86400.0)
    if isinstance(v, (int, float)):
        return None, repr(float(v)) if isinstance(v, float) else str(v)
    s = str(v)
    if s in ERRORS:
        return "e", s
    return "str", escape(s)


CELL_RE = re.compile(rb"<c\b[^>]*?/>|<c\b[^>]*?>.*?</c>", re.S)
# openpyxl emits an empty <v /> placeholder on every formula cell; it has to be
# removed, not appended to, or readers stop at the empty one and see nothing.
EXISTING_V_RE = re.compile(rb"<v\b[^>]*/>|<v\b[^>]*>.*?</v>", re.S)


def inject_sheet(xml: bytes, title, vals, stats):
    def fix(m):
        blob = m.group(0)
        if b"<f" not in blob:
            return blob
        head_end = blob.index(b">") + 1
        head, body = blob[:head_end], blob[head_end:]
        ref = re.search(rb'r="([A-Z]+\d+)"', head)
        if not ref:
            return blob
        coord = ref.group(1).decode()
        key = (title, coord)
        if key not in vals:
            # A formula whose result is an empty string reads back as None; Excel
            # still expects a string cell, so give it one.
            stats["empty"] += 1
            t, text = "str", ""
        else:
            t, text = render(vals[key])
            stats["filled"] += 1
        head = re.sub(rb'\st="[^"]*"', b"", head)
        if t:
            head = head[:-1] + f' t="{t}"'.encode() + b">"
        body = EXISTING_V_RE.sub(b"", body)
        body = body.replace(b"</c>", f"<v>{text}</v></c>".encode())
        return head + body

    return CELL_RE.sub(fix, xml)


def main():
    print(f"donor recalculation of {WORK.name} ...")
    shutil.copy(BUILT, WORK)
    proc = subprocess.run(["python3", RECALC, str(WORK), "1800"],
                          capture_output=True, text=True)
    print("  " + proc.stdout.strip().replace("\n", "\n  ")[:400])
    if '"error"' in proc.stdout:
        sys.exit("recalculation failed; nothing injected")

    vals = cached_values(WORK)
    print(f"harvested {len(vals):,} computed values")

    parts = sheet_parts(BUILT)
    stats = {"filled": 0, "empty": 0}
    tmp = SCRATCH / "injected.xlsx"
    with zipfile.ZipFile(BUILT) as src, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as dst:
        part_to_title = {v: k for k, v in parts.items()}
        for item in src.infolist():
            data = src.read(item.filename)
            title = part_to_title.get(item.filename)
            if title:
                data = inject_sheet(data, title, vals, stats)
            dst.writestr(item, data)
    shutil.move(tmp, BUILT)
    print(f"injected {stats['filled']:,} values ({stats['empty']:,} empty-string results)")


if __name__ == "__main__":
    main()
