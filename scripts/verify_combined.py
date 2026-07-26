#!/usr/bin/env python3
"""Independently re-derive the combined verdict from the ORIGINAL uploads and
check the delivered workbook agrees, cell by cell."""
import sys
from collections import Counter, defaultdict

from openpyxl import load_workbook

UPLOADS = "/root/.claude/uploads/217becff-8454-5f9a-a065-0284fc958152"
SRC = {
    "A": (f"{UPLOADS}/4475bf77-AFable__EasyHardMoney______NDX_______2026R4.5.3_2.xlsx", "Y", "AX"),
    "B": (f"{UPLOADS}/22b92ca6-B5.6_Sol_EasyHardMoney______NDX_______2026R4.5_SOL_CodexReviewed.xlsx", "R", "Q"),
    "C": (f"{UPLOADS}/0117f8a8-CGrok_EasyHardMoney______NDX_______2026R4.5.3_GROK.xlsx", "AH", "AG"),
}
OUT = "/home/user/Easy-Hard/EasyHardMoney_NDX_2026R4.5.3_COMBINED_ABC.xlsx"
FIRST, LAST = 5, 2519
ERRS = ("#DIV/0!", "#N/A", "#NAME?", "#NULL!", "#NUM!", "#REF!", "#VALUE!", "#ERROR!", "Err:")

fails = []


def check(label, ok, detail=""):
    print(f"  [{'PASS' if ok else 'FAIL'}] {label}{(' -- ' + detail) if detail else ''}")
    if not ok:
        fails.append(label)


# ---- 1. ground truth straight from the untouched uploads -----------------
print("== re-deriving ground truth from the original uploads ==")
truth = {}
src_raw = {}
dates = {}
scores = {}
for k, (path, zcol, scol) in SRC.items():
    wb = load_workbook(path, data_only=True)
    ws = wb["02A_每日分段"]
    truth[k] = [ws[f"{zcol}{r}"].value for r in range(FIRST, LAST + 1)]
    src_raw[k] = {c: [ws[f"{c}{r}"].value for r in range(FIRST, LAST + 1)]
                  for c in ("E", "F", "G", "H", "I", "K")}
    scores[k] = [ws[f"{scol}{r}"].value for r in range(FIRST, LAST + 1)]
    dates[k] = [ws[f"B{r}"].value for r in range(FIRST, LAST + 1)]
    wb.close()

n = LAST - FIRST + 1
expected = []
for i in range(n):
    za, zb, zc = truth["A"][i], truth["B"][i], truth["C"][i]
    expected.append("EASY" if za == zb == zc == "EASY"
                    else ("HARD" if za == zb == zc == "HARD" else "UNCERTAIN"))
exp_counts = Counter(expected)
print(f"  ground-truth combined counts: {dict(exp_counts)}  (n={n})")

# ---- 2. delivered workbook, cached values --------------------------------
print("\n== reading delivered workbook (cached values) ==")
wv = load_workbook(OUT, data_only=True)
wf = load_workbook(OUT)
print(f"  sheets ({len(wv.sheetnames)}): {wv.sheetnames}")

d = wv["01_整合判區"]
got = [d[f"N{r}"].value for r in range(FIRST, LAST + 1)]
got_a = [d[f"G{r}"].value for r in range(FIRST, LAST + 1)]
got_b = [d[f"I{r}"].value for r in range(FIRST, LAST + 1)]
got_c = [d[f"K{r}"].value for r in range(FIRST, LAST + 1)]
got_dates = [d[f"B{r}"].value for r in range(FIRST, LAST + 1)]
got_align = [d[f"P{r}"].value for r in range(FIRST, LAST + 1)]

# ---- 3. the rule itself ---------------------------------------------------
print("\n== rule 1-4: combined verdict ==")
bad = [(i, expected[i], got[i]) for i in range(n) if expected[i] != got[i]]
check("整合判區 matches ground truth on all 2515 days", not bad,
      f"{len(bad)} mismatches, first={bad[:3]}")
check("EASY count == all-three-EASY count", Counter(got)["EASY"] == exp_counts["EASY"],
      f"got {Counter(got)['EASY']} expected {exp_counts['EASY']}")
check("HARD count == all-three-HARD count", Counter(got)["HARD"] == exp_counts["HARD"],
      f"got {Counter(got)['HARD']} expected {exp_counts['HARD']}")
check("UNCERTAIN count == remainder", Counter(got)["UNCERTAIN"] == exp_counts["UNCERTAIN"],
      f"got {Counter(got)['UNCERTAIN']} expected {exp_counts['UNCERTAIN']}")
check("no verdict outside the three labels",
      set(got) <= {"EASY", "UNCERTAIN", "HARD"}, str(set(got)))

# every EASY day really is EASY in all three source files
easy_idx = [i for i in range(n) if got[i] == "EASY"]
check("every combined EASY day is EASY in A, B and C simultaneously",
      all(truth["A"][i] == truth["B"][i] == truth["C"][i] == "EASY" for i in easy_idx),
      f"{len(easy_idx)} EASY days checked")
hard_idx = [i for i in range(n) if got[i] == "HARD"]
check("every combined HARD day is HARD in A, B and C simultaneously",
      all(truth["A"][i] == truth["B"][i] == truth["C"][i] == "HARD" for i in hard_idx),
      f"{len(hard_idx)} HARD days checked")
# no missed day: any day where all three agree EASY must be combined EASY
missed_e = [i for i in range(n)
            if truth["A"][i] == truth["B"][i] == truth["C"][i] == "EASY" and got[i] != "EASY"]
missed_h = [i for i in range(n)
            if truth["A"][i] == truth["B"][i] == truth["C"][i] == "HARD" and got[i] != "HARD"]
check("no all-EASY day was wrongly demoted", not missed_e, f"{len(missed_e)} missed")
check("no all-HARD day was wrongly demoted", not missed_h, f"{len(missed_h)} missed")
unc_idx = [i for i in range(n) if got[i] == "UNCERTAIN"]
check("every UNCERTAIN day has at least one dissenting engine",
      all(not (truth["A"][i] == truth["B"][i] == truth["C"][i] in ("EASY", "HARD"))
          for i in unc_idx), f"{len(unc_idx)} UNCERTAIN days checked")

# ---- 4. rule 5: originals carried through --------------------------------
print("\n== rule 5: original A/B/C content carried into the combined file ==")
for k, (path, zcol, scol) in SRC.items():
    src = load_workbook(path)
    missing = [s for s in src.sheetnames if f"{k}_{s}" not in wv.sheetnames]
    check(f"{k}: all {len(src.sheetnames)} original sheets present as {k}_*", not missing, str(missing))
    src.close()

print("\n  -- source zone/score columns reproduced verbatim --")
for k, arr, col in (("A", got_a, "G"), ("B", got_b, "I"), ("C", got_c, "K")):
    bad = [i for i in range(n) if arr[i] != truth[k][i]]
    check(f"{k} 判區 column reproduces original {SRC[k][1]} column", not bad,
          f"{len(bad)} diffs, first={[(i, truth[k][i], arr[i]) for i in bad[:3]]}")
for k, col in (("A", "F"), ("B", "H"), ("C", "J")):
    arr = [d[f"{col}{r}"].value for r in range(FIRST, LAST + 1)]
    bad = [i for i in range(n)
           if not (arr[i] is None and scores[k][i] is None)
           and abs((arr[i] or 0) - (scores[k][i] or 0)) > 1e-9]
    check(f"{k} 分數 column reproduces original {SRC[k][2]} column", not bad,
          f"{len(bad)} diffs, first={[(i, scores[k][i], arr[i]) for i in bad[:3]]}")

print("\n  -- raw market data in the copied originals is byte-identical --")
for k, (path, _, _) in SRC.items():
    src = load_workbook(path, data_only=True)
    sws = src["02A_每日分段"]
    dws = wv[f"{k}_02A_每日分段"]
    diffs = 0
    checked = 0
    for r in range(FIRST, LAST + 1):
        for col in ("B", "E", "F", "G", "H", "I", "K"):   # date + raw market inputs
            a, b = sws[f"{col}{r}"].value, dws[f"{col}{r}"].value
            checked += 1
            if isinstance(a, float) and isinstance(b, float):
                if abs(a - b) > 1e-9:
                    diffs += 1
            elif a != b:
                diffs += 1
    check(f"{k}_02A raw inputs identical to upload", diffs == 0,
          f"{diffs} diffs of {checked} cells")
    src.close()

# ---- 5. alignment and integrity ------------------------------------------
print("\n== data alignment and formula integrity ==")
check("date column pulled from A matches original dates",
      all(got_dates[i] == dates["A"][i] for i in range(n)))
check("A/B/C dates agree on every row (source data)",
      all(dates["A"][i] == dates["B"][i] == dates["C"][i] for i in range(n)))
check("live 日期對齊 check column is OK on every row",
      set(got_align) == {"OK"}, str(Counter(got_align)))

err_cells = defaultdict(list)
formula_cells = 0
for ws in wf.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if c.data_type == "f":
                formula_cells += 1
for ws in wv.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and any(c.value.startswith(e) for e in ERRS):
                err_cells[c.value].append(f"{ws.title}!{c.coordinate}")
check("no formula error values anywhere in the workbook", not err_cells,
      str({k: (len(v), v[:3]) for k, v in err_cells.items()}))
print(f"  live formula cells in delivered file: {formula_cells:,}")

nones = [r for r in range(FIRST, LAST + 1) if d[f"N{r}"].value is None]
check("no blank verdict cell", not nones, f"{len(nones)} blanks")

# ---- 6. summary sheets agree with ground truth ---------------------------
print("\n== summary sheets ==")
s = wv["05_整合統計摘要"]
for i, zone in enumerate(("EASY", "UNCERTAIN", "HARD")):
    r = 8 + i
    check(f"05 摘要 {zone} count == {exp_counts[zone]}", s[f"B{r}"].value == exp_counts[zone],
          f"got {s[f'B{r}'].value}")
for j, (k, col) in enumerate((("A", "D"), ("B", "F"), ("C", "H"))):
    src_counts = Counter(truth[k])
    for i, zone in enumerate(("EASY", "UNCERTAIN", "HARD")):
        r = 8 + i
        check(f"05 摘要 {k} 引擎 {zone} == {src_counts[zone]}",
              s[f"{col}{r}"].value == src_counts[zone], f"got {s[f'{col}{r}'].value}")
check("05 摘要 total == 2515", s["B11"].value == n, f"got {s['B11'].value}")
check("05 摘要 日期對齊 status is OK", s["B5"].value == "全部對齊 OK", str(s["B5"].value))

yr = defaultdict(Counter)
for i in range(n):
    yr[dates["A"][i].year][expected[i]] += 1
YEAR_TOTAL_ROW = 35
for i, year in enumerate(range(2016, 2027)):
    r = 15 + i
    ok = (s[f"B{r}"].value == yr[year]["EASY"]
          and s[f"C{r}"].value == yr[year]["UNCERTAIN"]
          and s[f"D{r}"].value == yr[year]["HARD"])
    check(f"05 摘要 year {year} row", ok,
          f"got {(s[f'B{r}'].value, s[f'C{r}'].value, s[f'D{r}'].value)} "
          f"expected {(yr[year]['EASY'], yr[year]['UNCERTAIN'], yr[year]['HARD'])}")
check("05 摘要 year totals sum to 2515",
      s[f"E{YEAR_TOTAL_ROW}"].value == n, f"got {s[f'E{YEAR_TOTAL_ROW}'].value}")
check("05 摘要 year rows past the data are blank",
      all(s[f"A{r}"].value in (None, "") for r in range(26, YEAR_TOTAL_ROW)),
      str([s[f"A{r}"].value for r in range(26, YEAR_TOTAL_ROW)]))
check("05 摘要 coverage-check cell confirms full coverage",
      isinstance(s[f"B{YEAR_TOTAL_ROW+1}"].value, str)
      and "已涵蓋全部" in s[f"B{YEAR_TOTAL_ROW+1}"].value,
      str(s[f"B{YEAR_TOTAL_ROW+1}"].value))

# the three engines share six raw inputs; Q flags any day where they diverge
q = [d[f"Q{r}"].value for r in range(FIRST, LAST + 1)]
same = [i for i in range(n)
        if all(abs((src_raw[k][c][i] or 0) - (src_raw["A"][c][i] or 0)) < 1e-9
               if isinstance(src_raw["A"][c][i], (int, float))
               else src_raw[k][c][i] == src_raw["A"][c][i]
               for k in "BC" for c in ("E", "F", "G", "H", "I", "K"))]
check("輸入一致 column agrees with a direct comparison of the six shared inputs",
      [i for i in range(n) if q[i] == "一致"] == same,
      f"sheet says {q.count('一致')} 一致, direct comparison says {len(same)}")
check("每個 輸入不同 day is genuinely different in at least one input",
      all(i not in same for i in range(n) if q[i] == "輸入不同"),
      f"{q.count('輸入不同')} flagged")

# B ships an Excel Table on its review log; openpyxl drops it unless copied
src_b = load_workbook(SRC["B"][0])
st = src_b["10_審查與修訂紀錄"].tables
dt = wf["B_10_審查與修訂紀錄"].tables
check("B's Excel Table (ListObject) survived the merge",
      set(st) == set(dt) and all(st[k].ref == dt[k].ref
                                 and len(st[k].tableColumns) == len(dt[k].tableColumns)
                                 for k in st),
      f"source={ {k: st[k].ref for k in st} } delivered={ {k: dt[k].ref for k in dt} }")
src_b.close()

m = wv["06_三方比對"]
pairs = [("A", "B", 5), ("A", "C", 6), ("B", "C", 7)]
for x, y, r in pairs:
    exp = sum(1 for i in range(n) if truth[x][i] == truth[y][i])
    check(f"06 比對 {x}-{y} agreement == {exp}", m[f"B{r}"].value == exp, f"got {m[f'B{r}'].value}")
exp3 = sum(1 for i in range(n) if truth["A"][i] == truth["B"][i] == truth["C"][i])
check(f"06 比對 three-way agreement == {exp3}", m["B8"].value == exp3, f"got {m['B8'].value}")
evotes = Counter(sum(1 for k in "ABC" if truth[k][i] == "EASY") for i in range(n))
hvotes = Counter(sum(1 for k in "ABC" if truth[k][i] == "HARD") for i in range(n))
for v in range(4):
    r = 12 + v
    check(f"06 比對 EASY votes=={v} -> {evotes[v]}", m[f"B{r}"].value == evotes[v], f"got {m[f'B{r}'].value}")
    check(f"06 比對 HARD votes=={v} -> {hvotes[v]}", m[f"D{r}"].value == hvotes[v], f"got {m[f'D{r}'].value}")
check("06 比對: EASY-3-votes equals combined EASY count", m["B15"].value == exp_counts["EASY"])
check("06 比對: HARD-3-votes equals combined HARD count", m["D15"].value == exp_counts["HARD"])

# cross-tabs
ZS = ("EASY", "UNCERTAIN", "HARD")
for base, (x, y) in ((17, ("A", "B")), (24, ("A", "C")), (31, ("B", "C"))):
    tot = 0
    ok = True
    for i, zr in enumerate(ZS):
        for j, zc in enumerate(ZS):
            cell = m.cell(row=base + 2 + i, column=2 + j).value
            exp = sum(1 for t in range(n) if truth[x][t] == zr and truth[y][t] == zc)
            tot += cell or 0
            if cell != exp:
                ok = False
    check(f"06 交叉表 {x}x{y} all 9 cells correct and sum to {n}", ok and tot == n, f"sum={tot}")

# ---- 7. run table ---------------------------------------------------------
print("\n== 07_整合區間 run table ==")
runs = []
start = 0
for i in range(1, n + 1):
    if i == n or expected[i] != expected[start]:
        runs.append((expected[start], dates["A"][start], dates["A"][i - 1], i - start))
        start = i
p = wv["07_整合區間"]
check(f"run count == {len(runs)}", p[f"A{4 + len(runs)}"].value == len(runs)
      and p[f"A{5 + len(runs)}"].value is None, f"table rows for {len(runs)} runs")
bad = []
for i, (zone, d0, d1, days) in enumerate(runs):
    r = 5 + i
    if (p[f"B{r}"].value != zone or p[f"C{r}"].value != d0
            or p[f"D{r}"].value != d1 or p[f"E{r}"].value != days):
        bad.append((i + 1, zone, d0, d1, days,
                    p[f"B{r}"].value, p[f"C{r}"].value, p[f"D{r}"].value, p[f"E{r}"].value))
check("every run row (區制/起訖日/日數) matches ground truth", not bad,
      f"{len(bad)} bad, first={bad[:2]}")
tot_days = sum(p[f"E{5+i}"].value or 0 for i in range(len(runs)))
check(f"run 日數 sums to {n}", tot_days == n, f"got {tot_days}")

print("\n" + "=" * 72)
if fails:
    print(f"RESULT: {len(fails)} CHECK(S) FAILED")
    for f in fails:
        print("   -", f)
    sys.exit(1)
print("RESULT: ALL CHECKS PASSED")
