#!/usr/bin/env python3
"""Build the COMBINED workbook from source engines A / B / C.

Combination rule (user spec):
  EASY      only when A, B and C are all EASY that day
  HARD      only when A, B and C are all HARD that day
  UNCERTAIN every other day

Every original sheet of all three engines is carried into the output with an
A_ / B_ / C_ prefix, so the combined judgement stays traceable to live source
formulas rather than pasted numbers.
"""
import re
from copy import copy

from openpyxl import load_workbook, Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties

UPLOADS = "/root/.claude/uploads/217becff-8454-5f9a-a065-0284fc958152"
SRC = {
    "A": f"{UPLOADS}/4475bf77-AFable__EasyHardMoney______NDX_______2026R4.5.3_2.xlsx",
    "B": f"{UPLOADS}/22b92ca6-B5.6_Sol_EasyHardMoney______NDX_______2026R4.5_SOL_CodexReviewed.xlsx",
    "C": f"{UPLOADS}/0117f8a8-CGrok_EasyHardMoney______NDX_______2026R4.5.3_GROK.xlsx",
}
SRC_LABEL = {
    "A": "A = Fable  EasyHardMoney NDX 2026R4.5.3_2",
    "B": "B = Sol 5.6  EasyHardMoney NDX 2026R4.5 SOL CodexReviewed",
    "C": "C = Grok  EasyHardMoney NDX 2026R4.5.3 GROK",
}
# Where each engine keeps its verdict on sheet 02A_每日分段.
ZONE_COL = {"A": "Y", "B": "R", "C": "AH"}
SCORE_COL = {"A": "AX", "B": "Q", "C": "AG"}
DAILY = "02A_每日分段"

FIRST, LAST = 5, 2519          # data rows on every 02A sheet
RUN_ROWS = 400                 # capacity of the run table (301 runs today; N3 warns if exceeded)
OUT = "/home/user/Easy-Hard/EasyHardMoney_NDX_2026R4.5.3_COMBINED_ABC.xlsx"

# --- shared look, taken from the source workbooks -------------------------
NAVY = "1F4E79"
GREEN, YELLOW, RED = "C6EFCE", "FFEB9C", "FFC7CE"
F_TITLE = Font(name="Arial", sz=12, b=True, color=NAVY)
F_NOTE = Font(name="Arial", sz=9, color="555555")
F_HEAD = Font(name="Arial", sz=9, b=True, color="FFFFFF")
F_BODY = Font(name="Arial", sz=9)
F_BOLD = Font(name="Arial", sz=9, b=True)
F_KEY = Font(name="Arial", sz=9, b=True, color=NAVY)
FILL_HEAD = PatternFill("solid", fgColor=NAVY)
FILL_BAND = PatternFill("solid", fgColor="DDEBF7")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center")
WRAP = Alignment(vertical="top", wrap_text=True)


def zone_rules(ws, ref):
    """Green / yellow / red on a column of EASY-UNCERTAIN-HARD strings."""
    for text, colour in (("EASY", GREEN), ("UNCERTAIN", YELLOW), ("HARD", RED)):
        ws.conditional_formatting.add(
            ref,
            CellIsRule(operator="equal", formula=[f'"{text}"'],
                       fill=PatternFill("solid", bgColor=colour)),
        )


def head(ws, row, labels, widths=None):
    for i, text in enumerate(labels, start=1):
        if not text:          # spacer column -- leave it unpainted
            continue
        c = ws.cell(row=row, column=i, value=text)
        c.font, c.fill, c.alignment, c.border = F_HEAD, FILL_HEAD, CENTER, BOX
    if widths:
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w


NDAYS = None          # set once the sheet layout constants are known


def over_days(numerator):
    """numerator / trading-day count, blank rather than #DIV/0! on an empty table."""
    return f'=IF({NDAYS}=0,"",{numerator}/{NDAYS})'


def title(ws, text, note=None):
    c = ws.cell(row=1, column=1, value=text)
    c.font = F_TITLE
    if note:
        n = ws.cell(row=2, column=1, value=note)
        n.font = F_NOTE


# =========================================================================
# 1. Base workbook = engine A, sheets renamed A_*, cross-sheet refs rewritten
# =========================================================================
def rewrite_refs(formula, names, prefix):
    """Prefix every quoted sheet reference that names a sheet we renamed."""
    return re.sub(
        r"'([^']+)'!",
        lambda m: f"'{prefix}{m.group(1)}'!" if m.group(1) in names else m.group(0),
        formula,
    )


print("loading A ...")
wb = load_workbook(SRC["A"])
a_names = set(wb.sheetnames)
rewritten = 0
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if c.data_type == "f" and isinstance(c.value, str):
                new = rewrite_refs(c.value, a_names, "A_")
                if new != c.value:
                    c.value = new
                    rewritten += 1
for name in list(wb.sheetnames):
    wb[name].title = f"A_{name}"
print(f"  A: {len(a_names)} sheets renamed, {rewritten} formulas re-pointed")


# =========================================================================
# 2. Copy engines B and C in, cell by cell, with resolved styles
# =========================================================================
def copy_sheet(src_ws, dst_ws, names, prefix):
    n_formula = 0
    for row in src_ws.iter_rows():
        for c in row:
            if c.value is None and not c.has_style and c.comment is None:
                continue
            d = dst_ws.cell(row=c.row, column=c.column)
            v = c.value
            if c.data_type == "f" and isinstance(v, str):
                v = rewrite_refs(v, names, prefix)
                n_formula += 1
            d.value = v
            # B's review log quotes formulas as prose; openpyxl would promote any
            # string starting with "=" to a live formula, so pin the original type.
            if d.data_type != c.data_type:
                d.data_type = c.data_type
            # Copy the resolved style even when the source cell carries none: its
            # font then comes from B/C's own workbook default, which differs from
            # the merged workbook's default and would otherwise silently change.
            d.font = copy(c.font)
            d.fill = copy(c.fill)
            d.border = copy(c.border)
            d.alignment = copy(c.alignment)
            d.number_format = c.number_format
            d.protection = copy(c.protection)
            if c.comment is not None:
                d.comment = Comment(c.comment.text, c.comment.author or prefix.rstrip("_"))

    for letter, dim in src_ws.column_dimensions.items():
        # A <col> entry can span several columns; column_dimensions keys only the
        # first, so expand the span or the trailing columns lose their width.
        span = range(dim.min, dim.max + 1) if dim.min and dim.max else None
        letters = [get_column_letter(i) for i in span] if span else [letter]
        for lt in letters:
            n = dst_ws.column_dimensions[lt]
            n.width, n.hidden, n.bestFit = dim.width, dim.hidden, dim.bestFit
    for idx, dim in src_ws.row_dimensions.items():
        n = dst_ws.row_dimensions[idx]
        n.height, n.hidden = dim.height, dim.hidden
    for rng in src_ws.merged_cells.ranges:
        dst_ws.merge_cells(str(rng))
    for cf in src_ws.conditional_formatting:
        for rule in cf.rules:
            dst_ws.conditional_formatting.add(str(cf.sqref), copy(rule))
    if src_ws.freeze_panes:
        dst_ws.freeze_panes = src_ws.freeze_panes
    if src_ws.auto_filter.ref:
        dst_ws.auto_filter.ref = src_ws.auto_filter.ref
    dst_ws.sheet_view.showGridLines = src_ws.sheet_view.showGridLines
    return n_formula


for label in ("B", "C"):
    print(f"loading {label} ...")
    src = load_workbook(SRC[label])
    names = set(src.sheetnames)
    total = 0
    for ws in src.worksheets:
        total += copy_sheet(ws, wb.create_sheet(f"{label}_{ws.title}"), names, f"{label}_")
    print(f"  {label}: {len(names)} sheets copied, {total} formulas carried over")
    src.close()


# =========================================================================
# 3. 01_整合判區 — the combined daily verdict
# =========================================================================
print("building 01_整合判區 ...")
d = wb.create_sheet("01_整合判區")
title(
    d,
    "整合判區 (A∩B∩C) | 三引擎同時 EASY 才 EASY, 同時 HARD 才 HARD, 其餘一律 UNCERTAIN",
    "N欄=整合判區(唯一結論)。G/I/K欄為A/B/C三引擎原始判區, 由各自原表活算帶入; "
    "F/H/J欄為三引擎原始加權合成分。P欄逐日核對三表日期對齊。"
    "Q/R 為灰底輔助欄, 只供 07_整合區間 定位每段區間的起訖列, 不是分析數據。",
)
head(
    d, 4,
    ["交易日#", "日期", "年份", "週內", "NDX收盤",
     "A分數", "A判區", "B分數", "B判區", "C分數", "C判區",
     "EASY票數", "HARD票數", "整合判區", "三方一致?", "日期對齊",
     "區間起點(輔助)", "區間編號(輔助)"],
    [8, 12, 6, 5, 11, 8, 11, 8, 11, 8, 11, 9, 9, 12, 9, 9, 9, 9],
)
d.freeze_panes = "C5"
d.auto_filter.ref = f"A4:P{LAST}"      # the two helper columns stay out of the filter
for _c in ("Q4", "R4"):
    d[_c].fill = PatternFill("solid", fgColor="808080")

for r in range(FIRST, LAST + 1):
    f = {
        "A": f"='A_{DAILY}'!$A{r}",
        "B": f"='A_{DAILY}'!$B{r}",
        "C": f"='A_{DAILY}'!$C{r}",
        "D": f"='A_{DAILY}'!$D{r}",
        "E": f"='A_{DAILY}'!$E{r}",
        "F": f"='A_{DAILY}'!${SCORE_COL['A']}{r}",
        "G": f"='A_{DAILY}'!${ZONE_COL['A']}{r}",
        "H": f"='B_{DAILY}'!${SCORE_COL['B']}{r}",
        "I": f"='B_{DAILY}'!${ZONE_COL['B']}{r}",
        "J": f"='C_{DAILY}'!${SCORE_COL['C']}{r}",
        "K": f"='C_{DAILY}'!${ZONE_COL['C']}{r}",
        "L": f'=($G{r}="EASY")+($I{r}="EASY")+($K{r}="EASY")',
        "M": f'=($G{r}="HARD")+($I{r}="HARD")+($K{r}="HARD")',
        "N": f'=IF($L{r}=3,"EASY",IF($M{r}=3,"HARD","UNCERTAIN"))',
        "O": f'=IF(AND($G{r}=$I{r},$I{r}=$K{r}),"一致","分歧")',
        "P": (f"=IF(AND('A_{DAILY}'!$B{r}='B_{DAILY}'!$B{r},"
              f"'B_{DAILY}'!$B{r}='C_{DAILY}'!$B{r}),\"OK\",\"MISMATCH\")"),
        "Q": (f"=1" if r == FIRST else f'=IF($N{r}<>$N{r-1},1,0)'),
        "R": (f"=1" if r == FIRST else f"=$R{r-1}+$Q{r}"),
    }
    for col, formula in f.items():
        c = d[f"{col}{r}"]
        c.value = formula
        c.font = F_BODY
    d[f"B{r}"].number_format = "yyyy-mm-dd"
    d[f"E{r}"].number_format = "#,##0.00"
    for col in ("F", "H", "J"):
        d[f"{col}{r}"].number_format = "0.0"
    for col in ("D", "G", "I", "K", "L", "M", "N", "O", "P"):
        d[f"{col}{r}"].alignment = CENTER
    d[f"N{r}"].font = F_BOLD

for col in ("G", "I", "K", "N"):
    zone_rules(d, f"{col}{FIRST}:{col}{LAST}")
d.conditional_formatting.add(
    f"P{FIRST}:P{LAST}",
    CellIsRule(operator="equal", formula=['"MISMATCH"'],
               fill=PatternFill("solid", bgColor=RED)),
)
d.conditional_formatting.add(
    f"O{FIRST}:O{LAST}",
    CellIsRule(operator="equal", formula=['"分歧"'],
               fill=PatternFill("solid", bgColor="F2F2F2")),
)


# =========================================================================
# 4. 05_整合統計摘要
# =========================================================================
print("building 05_整合統計摘要 ...")
s = wb.create_sheet("05_整合統計摘要")
NRNG = f"'01_整合判區'!$N${FIRST}:$N${LAST}"
YRNG = f"'01_整合判區'!$C${FIRST}:$C${LAST}"
NDAYS = f"COUNT('01_整合判區'!$A${FIRST}:$A${LAST})"
title(s, "整合統計摘要 (全部為活公式, 隨來源表重算)",
      "整合判區 = A∩B∩C。三引擎判區分歧時一律歸 UNCERTAIN, 故整合 EASY/HARD 必然少於任一單一引擎。")
for i, w in enumerate([26, 12, 11, 12, 11, 12, 11, 12, 11], start=1):
    s.column_dimensions[get_column_letter(i)].width = w

s["A4"] = "涵蓋期間"
s["A4"].font = F_KEY
s["B4"] = f"=MIN('01_整合判區'!$B${FIRST}:$B${LAST})"
s["C4"] = f"=MAX('01_整合判區'!$B${FIRST}:$B${LAST})"
s["D4"] = f"=COUNT('01_整合判區'!$A${FIRST}:$A${LAST})&\" 個交易日\""
s["B4"].number_format = s["C4"].number_format = "yyyy-mm-dd"
for col in "BCD":
    s[f"{col}4"].font = F_BODY

s["A5"] = "日期對齊檢查 (三表逐日)"
s["A5"].font = F_KEY
s["B5"] = (f"=IF(COUNTIF('01_整合判區'!$P${FIRST}:$P${LAST},\"OK\")"
           f"=COUNTA('01_整合判區'!$P${FIRST}:$P${LAST}),\"全部對齊 OK\",\"有 MISMATCH\")")
s["B5"].font = F_BOLD

head(s, 7, ["判區", "整合 (A∩B∩C)", "佔比", "A 引擎", "佔比", "B 引擎", "佔比", "C 引擎", "佔比"])
for i, zone in enumerate(("EASY", "UNCERTAIN", "HARD")):
    r = 8 + i
    s[f"A{r}"] = zone
    s[f"A{r}"].font = F_BOLD
    s[f"A{r}"].alignment = CENTER
    s[f"B{r}"] = f'=COUNTIF({NRNG},$A{r})'
    s[f"C{r}"] = over_days(f"$B{r}")
    for j, col in enumerate("GIK"):
        vc = get_column_letter(4 + j * 2)
        pc = get_column_letter(5 + j * 2)
        s[f"{vc}{r}"] = f"=COUNTIF('01_整合判區'!${col}${FIRST}:${col}${LAST},$A{r})"
        s[f"{pc}{r}"] = over_days(f"${vc}{r}")
    for col in "BDFH":
        s[f"{col}{r}"].number_format = "#,##0"
    for col in "CEGI":
        s[f"{col}{r}"].number_format = "0.0%"
    for col in "BCDEFGHI":
        s[f"{col}{r}"].font = F_BODY
        s[f"{col}{r}"].border = BOX
    s[f"A{r}"].border = BOX
r = 11
s[f"A{r}"] = "合計"
s[f"A{r}"].font = F_BOLD
for col in "BDFH":
    s[f"{col}{r}"] = f"=SUM({col}8:{col}10)"
    s[f"{col}{r}"].number_format = "#,##0"
    s[f"{col}{r}"].font = F_BOLD

s["A13"] = "分年度整合判區日數"
s["A13"].font = F_KEY
head(s, 14, ["年份", "EASY", "UNCERTAIN", "HARD", "合計", "EASY%", "", "", ""])
for i, year in enumerate(range(2016, 2027)):
    r = 15 + i
    s[f"A{r}"] = year
    s[f"A{r}"].number_format = "0"
    for j, zone in enumerate(("EASY", "UNCERTAIN", "HARD")):
        col = get_column_letter(2 + j)
        s[f"{col}{r}"] = f'=COUNTIFS({YRNG},$A{r},{NRNG},"{zone}")'
        s[f"{col}{r}"].number_format = "#,##0"
    s[f"E{r}"] = f"=SUM($B{r}:$D{r})"
    s[f"E{r}"].number_format = "#,##0"
    s[f"F{r}"] = f'=IF($E{r}=0,"",$B{r}/$E{r})'
    s[f"F{r}"].number_format = "0.0%"
    for col in "ABCDEF":
        s[f"{col}{r}"].font = F_BODY
        s[f"{col}{r}"].alignment = CENTER
        s[f"{col}{r}"].border = BOX
r = 26
s[f"A{r}"] = "合計"
s[f"A{r}"].font = F_BOLD
for col in "BCDE":
    s[f"{col}{r}"] = f"=SUM({col}15:{col}25)"
    s[f"{col}{r}"].number_format = "#,##0"
    s[f"{col}{r}"].font = F_BOLD
    s[f"{col}{r}"].alignment = CENTER
s[f"F{r}"] = f'=IF($E{r}=0,"",$B{r}/$E{r})'
s[f"F{r}"].number_format = "0.0%"
s[f"F{r}"].font = F_BOLD
s[f"F{r}"].alignment = CENTER


# =========================================================================
# 5. 06_三方比對
# =========================================================================
print("building 06_三方比對 ...")
m = wb.create_sheet("06_三方比對")
title(m, "三引擎判區比對 (為何整合後 EASY/HARD 大幅收窄)",
      "整合規則要求三票全同, 因此任何一個引擎唱反調的日子都會落入 UNCERTAIN。")
for i, w in enumerate([26, 14, 12, 14, 12, 14, 12], start=1):
    m.column_dimensions[get_column_letter(i)].width = w

GR = f"'01_整合判區'!$G${FIRST}:$G${LAST}"
IR = f"'01_整合判區'!$I${FIRST}:$I${LAST}"
KR = f"'01_整合判區'!$K${FIRST}:$K${LAST}"
NDAYS = f"COUNT('01_整合判區'!$A${FIRST}:$A${LAST})"

head(m, 4, ["兩兩判區一致度", "一致日數", "佔比", "", "", "", ""])
for i, (name, x, y) in enumerate((("A 與 B", GR, IR), ("A 與 C", GR, KR), ("B 與 C", IR, KR))):
    r = 5 + i
    m[f"A{r}"] = name
    m[f"B{r}"] = f"=SUMPRODUCT(--({x}={y}))"
    m[f"C{r}"] = over_days(f"$B{r}")
    m[f"B{r}"].number_format = "#,##0"
    m[f"C{r}"].number_format = "0.0%"
    for col in "ABC":
        m[f"{col}{r}"].font = F_BODY
        m[f"{col}{r}"].border = BOX
r = 8
m[f"A{r}"] = "三方完全一致"
m[f"A{r}"].font = F_BOLD
m[f"B{r}"] = f"=SUMPRODUCT(--({GR}={IR}),--({IR}={KR}))"
m[f"C{r}"] = over_days(f"$B{r}")
m[f"B{r}"].number_format = "#,##0"
m[f"C{r}"].number_format = "0.0%"
for col in "ABC":
    m[f"{col}{r}"].font = F_BOLD
    m[f"{col}{r}"].border = BOX

m["A10"] = "投票分布 — 三引擎中恰有 N 票判為 EASY / HARD 的日數"
m["A10"].font = F_KEY
head(m, 11, ["票數", "EASY 日數", "佔比", "HARD 日數", "佔比", "", ""])
for v in range(4):
    r = 12 + v
    m[f"A{r}"] = v
    m[f"B{r}"] = f"=COUNTIF('01_整合判區'!$L${FIRST}:$L${LAST},$A{r})"
    m[f"C{r}"] = over_days(f"$B{r}")
    m[f"D{r}"] = f"=COUNTIF('01_整合判區'!$M${FIRST}:$M${LAST},$A{r})"
    m[f"E{r}"] = over_days(f"$D{r}")
    for col in "BD":
        m[f"{col}{r}"].number_format = "#,##0"
    for col in "CE":
        m[f"{col}{r}"].number_format = "0.0%"
    for col in "ABCDE":
        m[f"{col}{r}"].font = F_BOLD if v == 3 else F_BODY
        m[f"{col}{r}"].alignment = CENTER
        m[f"{col}{r}"].border = BOX
m["G12"] = "3 票 = 整合 EASY / 整合 HARD"
m["G12"].font = F_NOTE

ZONES = ("EASY", "UNCERTAIN", "HARD")
row = 17
for pair_name, x, y in (("A (列) × B (欄)", GR, IR),
                        ("A (列) × C (欄)", GR, KR),
                        ("B (列) × C (欄)", IR, KR)):
    m[f"A{row}"] = f"交叉表 {pair_name}"
    m[f"A{row}"].font = F_KEY
    head(m, row + 1, ["", *ZONES, "", "", ""])
    for i, zr in enumerate(ZONES):
        rr = row + 2 + i
        m[f"A{rr}"] = zr
        m[f"A{rr}"].font = F_BOLD
        m[f"A{rr}"].alignment = CENTER
        m[f"A{rr}"].border = BOX
        for j, zc in enumerate(ZONES):
            cc = get_column_letter(2 + j)
            m[f"{cc}{rr}"] = f'=SUMPRODUCT(--({x}="{zr}"),--({y}="{zc}"))'
            m[f"{cc}{rr}"].number_format = "#,##0"
            m[f"{cc}{rr}"].font = F_BODY
            m[f"{cc}{rr}"].alignment = CENTER
            m[f"{cc}{rr}"].border = BOX
            if zr == zc:
                m[f"{cc}{rr}"].fill = FILL_BAND
    row += 7


# =========================================================================
# 6. 07_整合區間 — contiguous runs of the combined verdict
# =========================================================================
print("building 07_整合區間 ...")
p = wb.create_sheet("07_整合區間")
title(p, "整合判區連續區間 (由 01_整合判區 R欄區間編號活算)",
      "每段列出區制、起訖日、交易日數與期間 NDX 報酬。顯示幾段由 N2 的實際區間數活算決定, "
      "多餘列自動留白; 若來源改動後段數超出表格容量, N3 會出現警告。")
head(p, 4, ["區間#", "區制", "起始日", "結束日", "交易日數", "起始NDX", "結束NDX",
            "期內報酬%(首收→末收)", "含首日報酬%(前收→末收)",
            "", "起始列(輔助)", "結束列(輔助)"],
     [8, 12, 12, 12, 10, 11, 11, 19, 20, 3, 12, 12])
p.freeze_panes = "A5"
NPX = f"'01_整合判區'!$E${FIRST}:$E${LAST}"
RID = f"'01_整合判區'!$R${FIRST}:$R${LAST}"
p["N1"] = "實際區間數"
p["N1"].font = F_KEY
p["N2"] = f"=MAX({RID})"
p["N2"].font = F_BOLD
p["N3"] = (f'=IF($N$2>{RUN_ROWS},"⚠ 區間數超出本表 {RUN_ROWS} 列容量, 請延長",'
           f'"表格容量 {RUN_ROWS} 列, 足夠")')
p["N3"].font = F_NOTE
p.column_dimensions["N"].width = 22
p.cell(row=3, column=1,
       value="H欄=判區成立後才進場的報酬(單日區間必為0%); I欄=含觸發當日跳空的報酬。"
             "K/L 為輔助欄, 每段只算一次 MATCH/COUNTIF 供左側各欄引用。").font = F_NOTE
for i in range(RUN_ROWS):
    r = 5 + i
    idx = i + 1
    # Rows past the live run count blank themselves out, so the table tracks the
    # real number of runs instead of being pinned to however many exist today.
    p[f"A{r}"] = f'=IF({idx}>$N$2,"",{idx})'
    g = f'IF($A{r}="","",'
    # Resolve each run's first/last row once, then reuse -- two full-column scans
    # per run instead of nine.
    p[f"K{r}"] = f"={g}MATCH($A{r},{RID},0))"
    p[f"L{r}"] = f"={g}$K{r}+COUNTIF({RID},$A{r})-1)"
    p[f"B{r}"] = f"={g}INDEX('01_整合判區'!$N${FIRST}:$N${LAST},$K{r}))"
    p[f"C{r}"] = f"={g}INDEX('01_整合判區'!$B${FIRST}:$B${LAST},$K{r}))"
    p[f"D{r}"] = f"={g}INDEX('01_整合判區'!$B${FIRST}:$B${LAST},$L{r}))"
    p[f"E{r}"] = f"={g}$L{r}-$K{r}+1)"
    p[f"F{r}"] = f"={g}INDEX({NPX},$K{r}))"
    p[f"G{r}"] = f"={g}INDEX({NPX},$L{r}))"
    # H measures from the regime's own first close, so a one-day regime is 0% by
    # construction; I measures from the close before it, capturing the move that
    # flipped the regime in the first place.
    p[f"H{r}"] = f'={g}IF($F{r}>0,$G{r}/$F{r}-1,""))'
    p[f"I{r}"] = f'={g}IF($K{r}>1,$G{r}/INDEX({NPX},$K{r}-1)-1,""))'
    p[f"C{r}"].number_format = p[f"D{r}"].number_format = "yyyy-mm-dd"
    p[f"E{r}"].number_format = "#,##0"
    p[f"F{r}"].number_format = p[f"G{r}"].number_format = "#,##0.00"
    p[f"H{r}"].number_format = p[f"I{r}"].number_format = "0.00%"
    for col in ("A", "B", "C", "D", "E", "F", "G", "H", "I", "K", "L"):
        p[f"{col}{r}"].font = F_BODY
    for col in ("A", "B", "C", "D", "E", "K", "L"):
        p[f"{col}{r}"].alignment = CENTER
zone_rules(p, f"B5:B{4 + RUN_ROWS}")


# =========================================================================
# 7. 00_README_整合
# =========================================================================
print("building 00_README_整合 ...")
rd = wb.create_sheet("00_README_整合")
rd.column_dimensions["A"].width = 22
rd.column_dimensions["B"].width = 118
rd.sheet_view.showGridLines = False
title(rd, "EASY / HARD MONEY 三引擎整合檔 (A ∩ B ∩ C) — NDX 2016-07-08 ~ 2026-07-10")

rows = [
    ("整合規則",
     "① 三個引擎 (A / B / C) 同日全部為 EASY → 整合判區 = EASY。"
     "② 三個引擎同日全部為 HARD → 整合判區 = HARD。"
     "③ 其餘所有情況 (含任何一個引擎分歧) → 整合判區 = UNCERTAIN。"),
    ("來源檔 A", SRC_LABEL["A"] + f"  → 判區取自 A_{DAILY}!{ZONE_COL['A']} 欄, 分數取自 {SCORE_COL['A']} 欄"),
    ("來源檔 B", SRC_LABEL["B"] + f"  → 判區取自 B_{DAILY}!{ZONE_COL['B']} 欄, 分數取自 {SCORE_COL['B']} 欄"),
    ("來源檔 C", SRC_LABEL["C"] + f"  → 判區取自 C_{DAILY}!{ZONE_COL['C']} 欄, 分數取自 {SCORE_COL['C']} 欄"),
    ("原檔完整保留",
     "三個來源檔的每一張工作表都原樣併入本檔, 分別加上 A_ / B_ / C_ 前綴 (含原始市場數據、"
     "指標庫權重、子分公式、統計摘要與監察面板)。跨表公式已同步改指到加前綴後的表名, 仍為活公式。"),
    ("三引擎各自門檻",
     "A: 合成分 ≥70 EASY / >40 UNCERTAIN / 其餘 HARD (18項權重)。 "
     "B: ≥92 EASY / ≥45 UNCERTAIN / 其餘 HARD (5項: T1=12 T3=28 B1=18 C1=14 V1=28)。 "
     "C: ≥80 EASY / ≥40 UNCERTAIN / 其餘 HARD (18項)。三者門檻與指標數不同, 這正是整合後大幅收窄的原因。"),
    ("資料對齊",
     "三個來源檔的 02A_每日分段 皆為同一組 2,515 個交易日, 逐列同日對齊 (第5列=2016-07-08, "
     "第2519列=2026-07-10)。01_整合判區 P 欄逐日活算核對三表日期是否相同, 全欄應為 OK。"),
    ("工作表",
     "00_README_整合 | 01_整合判區(主表, 2515日) | 05_整合統計摘要 | 06_三方比對 | 07_整合區間 | "
     "A_*(6張, Fable原檔) | B_*(9張, Sol原檔) | C_*(7張, Grok原檔)"),
    ("顏色",
     "綠=EASY (C6EFCE) | 黃=UNCERTAIN (FFEB9C) | 紅=HARD (FFC7CE)。三個來源檔本身對 UNCERTAIN "
     "的黃色並不一致 (B 與 C 用 FFEB9C, A 自己的 Y 欄用較深的 FFE699); 整合表採用三取二的 "
     "FFEB9C, 各來源原表則一律保留其原有顏色不動。"),
    ("哪些會自動重算",
     "四張整合表(01/05/06/07)全部為活公式, 不含寫死結果。來源方面: A 原檔帶 65,479 條活公式、"
     "B 原檔帶 23,634 條, 改動它們的輸入格會一路更新到整合判區; 但 C 原檔本身就只有數值、"
     "沒有任何公式(0條), 所以 C 判區欄是資料而非公式 — 要改 C 的判區須直接編輯 "
     "C_02A_每日分段 的 AH 欄。"),
    ("承襲自原檔的已知小問題",
     "A_01_指標庫 H20/H21 (T4=6、T5=26 兩項權重) 在來源檔 A 中就是以「文字」而非數字儲存。"
     "18 項權重按面值相加為 100, 但該表自己的合計格 H22 用的是 SUM(), 而 SUM() 會跳過文字, "
     "所以 H22 在來源檔 A 與本檔中都顯示 68 (=100-6-26)。A 的合成分公式是逐項相乘, 文字會被"
     "自動轉型, 因此判區結果不受影響 (已逐日核對與來源檔 A 完全一致)。此處按原樣保留未作更動; "
     "日後若要改動該權重表, 建議把 H20/H21 改回數值格式, H22 才會顯示 100。"),
]
r = 4
for k, v in rows:
    a, b = rd.cell(row=r, column=1, value=k), rd.cell(row=r, column=2, value=v)
    a.font, a.fill, a.alignment = F_KEY, FILL_BAND, Alignment(vertical="top")
    b.font, b.alignment = F_BODY, WRAP
    rd.row_dimensions[r].height = 30 if len(v) < 90 else 46
    r += 1

rd.cell(row=r + 1, column=1, value="整合結果 (活公式)").font = F_KEY
for i, zone in enumerate(ZONES):
    rr = r + 2 + i
    rd.cell(row=rr, column=1, value=zone).font = F_BOLD
    c = rd.cell(row=rr, column=2)
    c.value = (f'=IF({NDAYS}=0,"",COUNTIF({NRNG},$A{rr})&" 個交易日 ("'
               f'&TEXT(COUNTIF({NRNG},$A{rr})/{NDAYS},"0.0%")&")")')
    c.font = F_BODY


# =========================================================================
# 8. Sheet order: combined first, then A_*, B_*, C_*
# =========================================================================
# The A_* sheets arrive with page setup; give the new ones the same courtesy so
# the long tables print with a repeating header instead of naked columns.
for _name, _landscape, _repeat in (("00_README_整合", False, None), ("01_整合判區", True, "4:4"),
                                   ("05_整合統計摘要", False, None), ("06_三方比對", False, None),
                                   ("07_整合區間", True, "4:4")):
    _sh = wb[_name]
    _sh.page_setup.orientation = "landscape" if _landscape else "portrait"
    _sh.page_setup.fitToWidth, _sh.page_setup.fitToHeight = 1, 0
    _sh.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    if _repeat:
        _sh.print_title_rows = _repeat

front = ["00_README_整合", "01_整合判區", "05_整合統計摘要", "06_三方比對", "07_整合區間"]
order = front + [n for n in wb.sheetnames if n not in front]
wb._sheets = [wb[n] for n in order]
wb.active = 0
# Excel rebuilds every cached value on open, so the file is never stale.
wb.calculation.fullCalcOnLoad = True

print(f"saving {OUT} ...")
wb.save(OUT)
print("sheets:", wb.sheetnames)
print("done")
